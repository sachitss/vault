"""Languages: English, German, Nepali.

- dictionaries: same keys in de/ne, same {n} placeholders, no empty values
- first start follows the device language; the picker switches it and remembers it
- switching while unlocked redraws the screen and number formats
- German Excel/CSV export re-imports (column names matched in any language)
- reports come out in the chosen language
"""
import asyncio, json, pathlib, re, sys
from playwright.async_api import async_playwright

ROOT = pathlib.Path(__file__).resolve().parents[1]
URL = (ROOT / 'dist' / 'valuables-vault.html').resolve().as_uri()
OUT = ROOT / 'tests' / 'out'; OUT.mkdir(exist_ok=True)
PW = 'Language test passphrase 1!'
fail = []
def check(ok, msg):
    print(('ok   ' if ok else 'FAIL ') + msg)
    if not ok: fail.append(msg)

# dictionaries
de = json.loads((ROOT / 'src/i18n/de.json').read_text(encoding='utf-8'))
ne = json.loads((ROOT / 'src/i18n/ne.json').read_text(encoding='utf-8'))
check(set(de) == set(ne), f'de and ne have the same keys ({len(de)})')
ph = lambda s: sorted(set(re.findall(r'\{\d\}', s)))
bad = [k for d in (de, ne) for k, v in d.items() if not v or not set(ph(v)) <= set(ph(k))]
check(not bad, f'values non-empty, placeholders valid {bad[:3]}')


async def main():
    async with async_playwright() as p:
        b = await p.chromium.launch()
        # 1. German device -> German on first start
        ctx = await b.new_context(locale='de-DE', accept_downloads=True, viewport={'width': 1300, 'height': 900})
        pg = await ctx.new_page(); errs = []
        pg.on('pageerror', lambda e: errs.append(str(e)))
        await pg.goto(URL); await pg.wait_for_selector('#su-name')
        check(await pg.evaluate('LANG') == 'de', 'German device language detected')
        check('Verschlüsselten Tresor' in (await pg.text_content('#setup button[type=submit]')) or 'Tresor' in await pg.text_content('#setup'), 'setup screen in German')
        # 2. switch to Nepali on the setup screen
        await pg.click('.langpick [data-lang=ne]')
        check(await pg.evaluate("localStorage.getItem('vv-lang')") == 'ne', 'choice remembered')
        txt = await pg.text_content('#setup')
        check(re.search('[ऀ-ॿ]', txt) is not None, 'setup screen switched to Nepali')
        await pg.fill('#su-name', 'Test'); await pg.fill('#su-pw', PW); await pg.fill('#su-pw2', PW)
        await pg.click('#setup button[type=submit]'); await pg.wait_for_selector('#nav', timeout=60000)
        nav = await pg.text_content('#nav')
        check('ड्यासबोर्ड' in nav, 'navigation in Nepali')
        k = await pg.text_content('.kpi .v')
        # 3. switch to German while unlocked: screen redrawn, German number format
        await pg.evaluate("go('settings')")
        await pg.click('.langpick [data-lang=de]'); await pg.wait_for_timeout(200)
        check('Einstellungen' in await pg.text_content('#content h1'), 'settings redrawn in German')
        await pg.evaluate("go('dashboard')")
        kv = await pg.text_content('.kpi:nth-child(2) .v')
        check(re.search(r'\d\.\d{3}', kv) and '€' in kv, f'German number format on dashboard ({kv.strip()})')
        # 4. German exports: headers translated; CSV re-imports with all columns recognised
        await pg.evaluate("go('import')")
        async with pg.expect_download() as dl: await pg.click('#expf button[type=submit]')
        x = await dl.value; await x.save_as(OUT / 'export-de.xlsx')
        import openpyxl
        wb = openpyxl.load_workbook(OUT / 'export-de.xlsx')
        check('Detailliertes Inventar' in wb.sheetnames or any('Inventar' in n for n in wb.sheetnames), f'Excel sheet names in German {wb.sheetnames[:3]}')
        await pg.select_option('[name=fmt]', 'csv')
        async with pg.expect_download() as dl: await pg.click('#expf button[type=submit]')
        x = await dl.value; await x.save_as(OUT / 'export-de.csv')
        head = (OUT / 'export-de.csv').read_text(encoding='utf-8-sig').splitlines()[0]
        check('Inventar-ID' in head or 'Inventar' in head, f'CSV headers in German ({head[:60]}…)')
        await pg.set_input_files('#impfile', str(OUT / 'export-de.csv')); await pg.wait_for_timeout(500)
        mapped = await pg.evaluate("document.querySelectorAll('#impwiz select').length ? [...document.querySelectorAll('#impwiz select')].filter(s => s.value !== '' && s.value !== '-1').length : -1")
        check(mapped >= 20, f'German CSV columns matched on import ({mapped})')
        # 5. report in Nepali
        await pg.evaluate("setLang('ne')")
        rep = await pg.evaluate("new Promise(r => { openReport = (t, b) => r(reportHTML(t, b, true)); buildReport('insurance', activeItems(), {conf: false}); })")
        check('<html lang="ne"' in rep and 'बीमा' in rep and '@font-face' in rep, 'insurance report in Nepali with embedded font')
        # 6. lock screen keeps the language and offers the picker
        await pg.evaluate("lockVault(); renderLock()")
        check(await pg.locator('.lock-lang .langpick').count() == 1 and re.search('[ऀ-ॿ]', await pg.text_content('.lock')) is not None, 'lock screen in Nepali with picker')
        # 7. English device without a saved choice -> English, untouched
        ctx2 = await b.new_context(locale='en-US'); pg2 = await ctx2.new_page()
        await pg2.goto(URL); await pg2.wait_for_selector('#su-name')
        check(await pg2.evaluate('LANG') == 'en' and 'Create encrypted vault' in await pg2.text_content('#setup'), 'English device stays English')
        check(not errs, f'no page errors {errs[:2]}')
        await b.close()

asyncio.run(main())
print('\nAll i18n checks passed.' if not fail else f'\n{len(fail)} i18n checks FAILED')
sys.exit(1 if fail else 0)
