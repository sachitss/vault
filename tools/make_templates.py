import pathlib, os
ROOT = pathlib.Path(__file__).resolve().parents[1]
import json, csv
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment
S=json.load(open(ROOT/'src/schema.json', encoding='utf-8'))
# column list identical to the app (basic + unique detail fields)
det=[];seen=set()
for g,G in S['groups'].items():
  for f in G['fields']:
    if f['k'] in seen: continue
    seen.add(f['k']); det.append({**f,'label':f"{G['title']}: {f['label']}"})
cols=S['basicColumns']+det
labels=[c['label'] for c in cols]
def row_for(d):
  return [d.get(c['label'],'') for c in cols]
samples=[
 {'Inventory ID':'JWL-2026-00001','Item name':'Gold Necklace','Category code':'JWL','Subcategory':'Necklace','Description':'18K yellow gold curb-link necklace with pendant, 50 cm','Manufacturer':'Sample Goldschmiede','Date acquired':date(2021,6,15),'Purchase date':date(2021,6,15),'Purchase location':'Cologne','Seller / dealer':'Sample Juwelier GmbH','Country of purchase':'Germany','Invoice number':'INV-2021-12345','Purchase price':4500,'Currency':'EUR','Owner':'Sachit Shrestha','Beneficiaries (; separated)':'Child (sample)','Ownership %':100,'Ownership type':'Sole','Acquisition method':'Purchased','Storage location':'Bank locker – Sparkasse (sample)','Insurance policy (insurer)':'Sample Hausrat AG','Current estimated value':5900,'Current value date':date(2026,5,20),'Current value source':'Appraised – Professional valuer','Insurance value':7000,'Insurance value date':date(2026,5,20),'Metal & Gold: Metal type':'Gold','Metal & Gold: Gold type / colour':'Yellow','Metal & Gold: Purity (‰ fineness)':750,'Metal & Gold: Karat':'750 / 18K','Metal & Gold: Gross weight (g)':52.4,'Metal & Gold: Net metal weight (g)':51.8,'Metal & Gold: Hallmark / stamp':"750, maker's mark"},
 {'Inventory ID':'DIA-2026-00001','Item name':'Diamond solitaire ring','Category code':'DIA','Description':'Platinum solitaire ring, round brilliant 1.02 ct','Purchase date':date(2019,3,8),'Seller / dealer':'Sample Diamonds Antwerp','Country of purchase':'Belgium','Invoice number':'SD-19-0442','Purchase price':8200,'Currency':'EUR','Owner':'Family member (sample)','Ownership type':'Sole','Ownership %':100,'Acquisition method':'Purchased','Storage location':'Bedroom safe','Insurance policy (insurer)':'Sample Hausrat AG','Current estimated value':7400,'Current value date':date(2025,11,2),'Current value source':'Market – Dealer quote','Insurance value':11500,'Insurance value date':date(2025,11,2),'Diamond / Gemstone: Gemstone type':'Diamond','Diamond / Gemstone: Carat (total)':1.02,'Diamond / Gemstone: Shape':'Round','Diamond / Gemstone: Cut grade':'Excellent','Diamond / Gemstone: Colour':'G','Diamond / Gemstone: Clarity':'VS1','Diamond / Gemstone: Certification authority':'GIA','Diamond / Gemstone: Certificate number':'2195730461'},
 {'Inventory ID':'COI-2026-00001','Item name':'Gold Asarphi coins (5 pcs)','Category code':'COI','Description':'Family gold coins, 10 g each','Date acquired':date(2012,10,24),'Currency':'NPR','Owner':'Sachit Shrestha','Co-owners (; separated)':'Family member (sample)','Beneficiaries (; separated)':'Child (sample)','Ownership %':50,'Ownership type':'Family','Acquisition method':'Inherited','Storage location':'Bank locker – Kathmandu (sample)','Current estimated value':600000,'Current value date':date(2022,1,15),'Current value source':'Estimated – Owner estimate','Coin / Medal: Issuing country':'Nepal','Coin / Medal: Quantity':5,'Metal & Gold: Metal type':'Gold','Metal & Gold: Purity (‰ fineness)':999,'Metal & Gold: Net metal weight (g)':50},
 {'Inventory ID':'ART-2026-00001','Item name':'Himalayan landscape, oil on canvas','Category code':'ART','Description':'Oil on canvas, 80 × 60 cm, signed lower right','Purchase date':date(2023,9,30),'Seller / dealer':'Sample Gallery Kathmandu','Country of purchase':'Nepal','Purchase price':2800,'Currency':'EUR','Owner':'Sachit Shrestha','Acquisition method':'Purchased','Storage location':'Home – living room','Insurance policy (insurer)':'Sample Hausrat AG','Current estimated value':3200,'Current value date':date(2025,12,1),'Current value source':'Estimated – Owner estimate','Artwork: Artist':'Sample Artist','Artwork: Title of work':'Evening over Annapurna','Artwork: Creation year':'2021','Artwork: Medium':'Oil on canvas','Artwork: Signature':'Signed front'},
 {'Inventory ID':'WAT-2026-00001','Item name':'Automatic wristwatch','Category code':'WAT','Description':'Stainless steel, 40 mm','Brand':'Sample Watch Co.','Model':'Explorer 40','Serial number':'SWC-88214','Purchase date':date(2022,12,20),'Purchase price':6900,'Currency':'EUR','Owner':'Sachit Shrestha','Acquisition method':'Purchased','Watch: Reference no.':'214270','Watch: Box & papers':'Box and papers','Metal & Gold: Metal type':'Steel'},
]
# ---------- sample CSV (semicolon, ISO dates, UTF-8 BOM) ----------
def fmt(v): return v.isoformat() if isinstance(v,date) else v
with open(ROOT/'templates/sample-inventory.csv','w',newline='',encoding='utf-8-sig') as f:
  w=csv.writer(f,delimiter=';'); w.writerow(labels); [w.writerow([fmt(x) for x in row_for(s)]) for s in samples]
with open(ROOT/'templates/inventory-template.csv','w',newline='',encoding='utf-8-sig') as f:
  w=csv.writer(f,delimiter=';'); w.writerow(labels); w.writerow([fmt(x) for x in row_for(samples[0])])
# ---------- Excel template ----------
NAVY='1F3A5F'; F='Arial'
hdr=Font(name=F,bold=True,color='FFFFFF',size=10); hfill=PatternFill('solid',fgColor=NAVY)
reqfill=PatternFill('solid',fgColor='B23A3A'); inp=PatternFill('solid',fgColor='FFF7D6'); ex=Font(name=F,italic=True,color='7B8496',size=10)
base=Font(name=F,size=10); thin=Side(style='thin',color='E4E0D8')
wb=Workbook()
rm=wb.active; rm.title='Read me'
lines=[('Valuables Inventory – Import Template',Font(name=F,size=16,bold=True,color=NAVY)),('',None),
('How to use',Font(name=F,bold=True,size=11)),
('1. Fill one row per valuable in the sheet "Inventory Import". Row 2 is an example – overwrite or delete it.',base),
('2. Columns with a red header are required (Item name, Category code). Everything else is optional.',base),
('3. Yellow cells are input cells. Drop-down lists are provided for category, currency, acquisition and ownership type.',base),
('4. Dates: real Excel dates, or text as YYYY-MM-DD or DD.MM.YYYY. Numbers: 4500 or 4.500,00 both work.',base),
('5. Owners, storage locations and insurance policies are matched by name; new names are created on import.',base),
('6. Leave "Inventory ID" empty for new items – the app assigns CAT-YYYY-NNNNN. Fill it only to update existing records.',base),
('7. In Valuables Vault: Import / Export → Import → choose this file → check the column mapping → Validate & preview → Import.',base),
('',None),('Category codes',Font(name=F,bold=True,size=11))]
for t,f in lines: rm.append([t]); f and setattr(rm.cell(rm.max_row,1),'font',f)
for k,v in S['categories'].items(): rm.append([k,v['name'],'Certificate expected' if v['certExpected'] else '']); [setattr(c,'font',base) for c in rm[rm.max_row]]
rm.append([]); rm.append(['Value terms']); rm.cell(rm.max_row,1).font=Font(name=F,bold=True,size=11)
for a,b in [('Purchase price','What you paid (historic cost)'),('Current estimated value','What it would sell for today (market / appraised / estimate)'),('Insurance value','Replacement value – the cost of an equivalent new piece; usually higher than market value')]:
  rm.append([a,b]); [setattr(c,'font',base) for c in rm[rm.max_row]]
rm.append([]); rm.append(['Confidential: do not put locker numbers, policy numbers or ID numbers in this file unless it is stored encrypted.']); rm.cell(rm.max_row,1).font=Font(name=F,bold=True,color='B23A3A')
rm.column_dimensions['A'].width=26; rm.column_dimensions['B'].width=70; rm.column_dimensions['C'].width=22
# lists
ls=wb.create_sheet('Lists')
lists={'cat':list(S['categories'].keys()),'cur':S['currencies'],'acq':S['acquisition'],'own':S['ownershipTypes'],'loc':S['locationTypes'],'val':S['valuationTypes']}
for i,(k,v) in enumerate(lists.items(),1):
  ls.cell(1,i,k).font=Font(name=F,bold=True)
  for j,x in enumerate(v,2): ls.cell(j,i,x).font=base
ls.sheet_state='hidden'
def rng(k): i=list(lists).index(k)+1; return f"Lists!${L(i)}$2:${L(i)}${len(lists[k])+1}"
# inventory import
ws=wb.create_sheet('Inventory Import',1)
N=500
ws.append(labels)
for i,c in enumerate(cols,1):
  cell=ws.cell(1,i); cell.font=hdr; cell.fill=reqfill if c.get('required') else hfill; cell.alignment=Alignment(wrap_text=True,vertical='center')
  if c.get('hint'): cell.comment=Comment(c['hint'],'Template')
  ws.column_dimensions[L(i)].width=max(12,min(34,len(c['label'])+2))
ws.row_dimensions[1].height=42
ws.append(row_for(samples[0]))
for i,c in enumerate(cols,1):
  ce=ws.cell(2,i); ce.font=ex
  for r in range(2,N+2):
    x=ws.cell(r,i); x.fill=inp; x.border=Border(bottom=thin)
    if r>2: x.font=base
    if c.get('type')=='date': x.number_format='DD.MM.YYYY'
    elif c.get('type')=='number': x.number_format='#,##0.00' if any(w in c['label'].lower() for w in ['price','value']) else 'General'
ws.freeze_panes='C2'; ws.auto_filter.ref=f"A1:{L(len(cols))}1"
def dv(colkey,listkey,title):
  i=[c['k'] for c in cols].index(colkey)+1
  d=DataValidation(type='list',formula1=f'={rng(listkey)}',allow_blank=True,showErrorMessage=True,errorTitle=title,error=f'Choose a value from the list ({title}).'); ws.add_data_validation(d); d.add(f'{L(i)}2:{L(i)}{N+1}')
dv('cat','cat','Category code'); dv('currency','cur','Currency'); dv('acquisition','acq','Acquisition method'); dv('ownershipType','own','Ownership type')
for k in ['purchasePrice','currentValue','insuranceValue']:
  i=[c['k'] for c in cols].index(k)+1; d=DataValidation(type='decimal',operator='greaterThanOrEqual',formula1='0',allow_blank=True,showErrorMessage=True,error='Enter a number ≥ 0'); ws.add_data_validation(d); d.add(f'{L(i)}2:{L(i)}{N+1}')
i=[c['k'] for c in cols].index('ownershipPct')+1; d=DataValidation(type='decimal',operator='between',formula1='0',formula2='100',allow_blank=True,showErrorMessage=True,error='0–100'); ws.add_data_validation(d); d.add(f'{L(i)}2:{L(i)}{N+1}')
# summary with live formulas
sm=wb.create_sheet('Summary',2)
ci={c['k']:L(i) for i,c in enumerate(cols,1)}
sm.append(['Summary of "Inventory Import" (updates as you type; mixed currencies are not converted)']); sm['A1'].font=Font(name=F,bold=True,size=12,color=NAVY)
sm.append([]); sm.append(['Measure','Value']); [setattr(c,'font',hdr) or setattr(c,'fill',hfill) for c in sm[3]]
rows=[('Items entered',f"=COUNTA('Inventory Import'!{ci['name']}2:{ci['name']}{N+1})"),
('Total purchase price',f"=SUM('Inventory Import'!{ci['purchasePrice']}2:{ci['purchasePrice']}{N+1})"),
('Total current estimated value',f"=SUM('Inventory Import'!{ci['currentValue']}2:{ci['currentValue']}{N+1})"),
('Total insurance value',f"=SUM('Inventory Import'!{ci['insuranceValue']}2:{ci['insuranceValue']}{N+1})"),
('Items without owner',f"=SUMPRODUCT(('Inventory Import'!{ci['name']}2:{ci['name']}{N+1}<>\"\")*('Inventory Import'!{ci['owner']}2:{ci['owner']}{N+1}=\"\"))"),
('Items without storage location',f"=SUMPRODUCT(('Inventory Import'!{ci['name']}2:{ci['name']}{N+1}<>\"\")*('Inventory Import'!{ci['location']}2:{ci['location']}{N+1}=\"\"))"),
('Items without insurance value',f"=SUMPRODUCT(('Inventory Import'!{ci['name']}2:{ci['name']}{N+1}<>\"\")*('Inventory Import'!{ci['insuranceValue']}2:{ci['insuranceValue']}{N+1}=\"\"))")]
for a,b in rows: sm.append([a,b]); sm.cell(sm.max_row,1).font=base; sm.cell(sm.max_row,2).font=base; sm.cell(sm.max_row,2).number_format='#,##0'
sm.append([]); sm.append(['Values in mixed currencies are summed as entered. The app converts with your own exchange rates.']); sm.cell(sm.max_row,1).font=ex
sm.column_dimensions['A'].width=36; sm.column_dimensions['B'].width=18
# reference sheets (planning lists)
def ref(name,heads,example,note,dvs={}):
  s=wb.create_sheet(name); s.append([note]); s['A1'].font=ex; s.append(heads)
  for i in range(1,len(heads)+1): c=s.cell(2,i); c.font=hdr; c.fill=hfill; c.alignment=Alignment(wrap_text=True); s.column_dimensions[L(i)].width=max(13,len(heads[i-1])+3)
  s.append(example); [setattr(c,'font',ex) for c in s[3]]
  for r in range(3,103):
    for i in range(1,len(heads)+1): s.cell(r,i).fill=inp
  for col,k in dvs.items(): d=DataValidation(type='list',formula1=f'={rng(k)}',allow_blank=True); s.add_data_validation(d); d.add(f'{col}3:{col}102')
  s.freeze_panes='A3'; return s
v=ref('Valuation History',['Inventory ID','Date','Valuation type','Value','Currency','Source','Appraiser / valuer','Document'],['JWL-2026-00001',date(2026,5,20),'Insurance',7000,'EUR','Insurance appraisal','Sample Gutachter','appraisal-2026-05-20.pdf'],'Record additional valuations here. The prototype imports current and insurance values from the Inventory sheet; enter further history in the app.',{'C':'val','E':'cur'})
for r in range(3,103): v.cell(r,2).number_format='DD.MM.YYYY'; v.cell(r,4).number_format='#,##0.00'
ref('Storage Locations',['Location name','Type','Address / description','Certified safe (yes/no)','Bank','Branch','Locker no. (CONFIDENTIAL)','Holders','Bank-provided cover','Last inspection','Next review'],['Bank locker – Sparkasse (sample)','Bank locker','','','Sample Sparkasse','Main branch','','Sachit Shrestha; Family member (sample)',10000,date(2026,5,20),date(2027,5,20)],'Locations are created by name during import; complete locker and safe details in the app (Locations), where confidential numbers are encrypted.',{'B':'loc'})
ref('Insurance Policies',['Insurer','Policy type','Sum insured','Currency','Renewal date','Deductible','Valuables limit % of sum insured','Valuables outside safe limit','Jewellery limit','Gold / coins limit','Art limit','Per-item limit','Locker contents covered (yes/no)','Locker limit'],['Sample Hausrat AG','Household contents (Hausrat)',95000,'EUR',date(2026,11,22),250,20,20000,'','','','','yes',25000],'Policy limits as stated in your policy schedule. Enter them in the app (Insurance) for the automatic coverage analysis.',{'D':'cur'})
wb.move_sheet('Lists', offset=10)
wb.save(ROOT/'templates/valuables-import-template.xlsx'); print('ok',len(cols),'columns')
